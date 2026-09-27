"""v5.40 — Slides & Reports: a designed PowerPoint or a full report from one sentence.
The AI writes a plan, then each slide / section (canned answers here); code designs, draws,
checks figures and writes the .pptx / .docx / PDF.

  python3 tests/e2e_v540.py
"""
import base64, io, json, os, sys, zipfile
from playwright.sync_api import sync_playwright
from harness import Env, new_page, check, real_errors, finish, HERE

env = Env(); errors = []
SHOTS = os.environ.get("SHOTS")   # a folder: save screenshots there (for looking at the design)

NATIVE = """(() => { const N = window.AttuneNative, S = window.__mock;
  const R = (id, v) => setTimeout(() => window.__attuneNative.resolve(id, JSON.stringify(v)), 30);
  N.makePdf = (id, arg) => { (S.pdfCalls = S.pdfCalls || []).push(JSON.parse(arg)); R(id, { b64: btoa("%PDF-1.7 fake"), bytes: 13 }); };
  N.search = (id, arg) => { const a = JSON.parse(arg); (S.searches = S.searches || []).push(a.q);
    R(id, { via: "duckduckgo", hits: [
      { title: "Crane market Egypt 2026", url: "https://example.com/cranes-" + S.searches.length, text: "The Egyptian crane rental market reached EGP 4.2 billion in 2025. Mobile cranes are 61% of rentals. The New Capital has 1,350 active sites. Day rates: 60 t EGP 35,000, 200 t EGP 90,000, 500 t EGP 210,000." }] }); };
})();"""

def shot(page, name, sel=None):
    if not SHOTS: return
    (page.locator(sel) if sel else page).screenshot(path=os.path.join(SHOTS, name + ".png"))

def queue(page, answers):
    page.evaluate("(a) => { const M = window.__mock; M.bodies = []; M.fakeQueue = a; }", answers)

def sent(page):
    return page.evaluate("window.__mock.bodies.filter(b => b.max_tokens > 2).map(b => b.messages.map(m => m.content).join('\\n'))")

def canvas_drawn(page, sel):
    return page.evaluate("""(sel) => { const c = document.querySelector(sel); if (!c) return 0;
      const d = c.getContext('2d').getImageData(0, 0, c.width, c.height).data; const seen = new Set();
      for (let i = 0; i < d.length; i += 4 * 97) seen.add(d[i] + ',' + d[i + 1] + ',' + d[i + 2]); return seen.size; }""", sel)

def pptx_texts(b64):
    z = zipfile.ZipFile(io.BytesIO(base64.b64decode(b64)))
    slides = sorted([n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")], key=lambda n: int(n[16:-4]))
    return z, [z.read(n).decode() for n in slides]

OUTLINE = """TITLE: Adrighem & Aldibiki Cranes
SUBTITLE: Heavy lifting partner for the New Capital
1. [bullets] Who we are
2. [two] Renting vs buying a crane
3. [steps] How a lift is planned
4. [table] Our crane classes
5. [bullets] Why choose us"""
SLIDES = [
    "- Experience: 30 years lifting in Egypt\n- Fleet: Liebherr, Demag, Grove and more\n- Team: 150 trained people\nNOTES: We are a family company with a big fleet.",
    "LEFT: Renting\n- No capital tied up\n- Operator included\n- Maintenance is ours\nRIGHT: Buying\n- Big upfront cost\n- Idle months cost money\n- You train operators\nNOTES: Most contractors save by renting.",
    "- Site visit: we measure ground and radius\n- Lift plan: load chart and rigging checked\n- Permit: authorities informed\n- Lift day: a supervisor on site\nNOTES: Four steps, every time.",
    "Class | Typical job | Example\n20–60 t | Buildings | Grove GMK\n60–200 t | Bridges | Liebherr LTM 1100\n200–500 t | Wind turbines | Demag AC 500\nNOTES: Pick the class by the heaviest lift.",
    "- Safety: every lift planned in writing\n- Speed: a crane on site within 24 hours\n- Price: fair day rates, no hidden fees\nNOTES: Close with the offer.",
]

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx, page = new_page(br, env, errors)
    page.goto(env.url); page.wait_for_selector("nav", timeout=15000)
    page.locator("header button:has-text('No model')").click()
    page.locator("text=Recommended for this device").locator("xpath=..").get_by_role("button", name="Install").first.click()
    page.wait_for_selector("text=Running now", timeout=10000)
    page.evaluate("window.__attuneBack()"); page.wait_for_timeout(200)
    page.evaluate(NATIVE)
    page.locator("nav button").last.click(); page.wait_for_timeout(250)
    check(page.locator(".rounded-t-2xl button:has-text('Slides & Reports')").count() >= 1, "More lists Slides & Reports")
    page.locator(".rounded-t-2xl button:has-text('Slides & Reports')").first.click()
    page.wait_for_selector("[data-testid=slides]", timeout=5000)
    shot(page, "01-form")

    # ---- 1. a presentation from one sentence ----
    page.fill("[data-testid=slides-prompt]", "A pitch for our crane rental company to New Capital contractors")
    page.fill("[data-testid=slides-audience]", "contractors")
    page.locator("[data-testid=slides-n-5]").click()
    page.locator("[data-testid=slides-theme-steel]").click()
    queue(page, [OUTLINE] + SLIDES)
    page.locator("[data-testid=slides-go]").click()
    page.wait_for_selector("[data-testid=slides-save-pptx]", timeout=30000)
    msgs = sent(page)
    check(len(msgs) == 6, "one plan + one call per slide (%d calls)" % len(msgs))
    check("exactly 5 lines" in msgs[0] and "stats (" not in msgs[0] and "chart (" not in msgs[0], "the plan asks for 5 slides; no number / chart slides without sources")
    check("slide 2 of 5" in msgs[2] and '"Renting vs buying a crane"' in msgs[2] and "LEFT:" in msgs[2] and "Don't invent statistics" in msgs[2], "each slide is written on its own, in its format, told not to invent figures")
    check(page.locator("[data-testid=slides-deck-title]").inner_text() == "Adrighem & Aldibiki Cranes", "the deck's title")
    n = page.locator("[data-testid^=slide-canvas-]").count()
    check(n == 8, "cover + agenda + 5 slides + closing = 8 previews (%d)" % n)
    check(canvas_drawn(page, "[data-testid=slide-canvas-4]") > 3, "slides are drawn on the phone (preview canvas has content)")
    note = page.locator("[data-testid=slides-check-note]")
    check(note.count() == 1 and "AI's memory" in note.inner_text(), "figures without sources are flagged (\"30 years\", \"150\", \"24 hours\")")
    shot(page, "02-deck", "[data-testid=slides-deck]")
    for k in (1, 2, 3, 4, 5, 6, 7, 8):
        shot(page, "slide-%d" % k, "[data-testid=slide-canvas-%d]" % k)

    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    saved = page.evaluate("window.__mock.lastSaved")
    check(saved["name"] == "Adrighem & Aldibiki Cranes.pptx" and saved["mime"].endswith("presentationml.presentation"), "saved as a PowerPoint file (%s)" % saved["name"])
    z, xs = pptx_texts(saved["b64"])
    check(len(xs) == 8, "the .pptx has 8 slides")
    check("Renting" in xs[3] and "Buying" in xs[3] and "Liebherr LTM 1100" in xs[5] and "<a:tbl>" in xs[5], "comparison and table slides are in the file (a real PowerPoint table)")
    check("FACC15" in z.read("ppt/theme/theme1.xml").decode(), "the Construction design's colours are the file's theme")
    check("Most contractors save by renting." in z.read("ppt/notesSlides/notesSlide4.xml").decode(), "speaker notes are in the file")
    try:
        from pptx import Presentation
        pr = Presentation(io.BytesIO(base64.b64decode(saved["b64"])))
        check(len(pr.slides) == 8 and pr.slide_width == 12192000, "python-pptx opens it: 8 slides, 16:9")
    except ImportError:
        print("SKIP python-pptx not installed")

    # ---- 2. edit a slide by hand, then have the AI redo another ----
    page.locator("[data-testid=slide-thumb-3]").click()
    page.wait_for_selector("[data-testid=slide-editor]", timeout=5000)
    check("- Experience: 30 years lifting in Egypt" in page.input_value("[data-testid=slide-edit-body]"), "the editor shows the slide as simple lines")
    page.fill("[data-testid=slide-edit-title]", "Who we are — since 1995")
    page.fill("[data-testid=slide-edit-body]", "- Experience: 30 years lifting in Egypt\n- Fleet: 40 cranes, 20 to 500 t\n- Team: 150 trained people\n- Safety: zero lost-time accidents")
    shot(page, "03-editor")
    page.locator("[data-testid=slide-edit-save]").click()
    page.wait_for_selector("[data-testid=slides-deck]", timeout=5000)
    page.locator("[data-testid=slide-thumb-7]").click()
    page.wait_for_selector("[data-testid=slide-editor]", timeout=5000)
    queue(page, ["QUOTE: Every lift is planned, or it doesn't happen.\nBY:\nNOTES: Our promise."])
    page.fill("[data-testid=slide-ai-ask]", "one strong line")
    page.locator("[data-testid=slide-kind-quote]").click()
    page.wait_for_function("document.querySelector('[data-testid=slide-edit-body]').value.includes('QUOTE: Every lift')", timeout=15000)
    m = sent(page)
    check(len(m) == 1 and "QUOTE:" in m[0] and "ALSO: one strong line" in m[0], "“Turn it into a quote” asks the AI for that slide only, with the request")
    page.locator("[data-testid=slide-edit-save]").click()
    page.wait_for_selector("[data-testid=slides-deck]", timeout=5000)
    page.locator("[data-testid=deck-theme-ocean]").click(); page.wait_for_timeout(200)
    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    z, xs = pptx_texts(page.evaluate("window.__mock.lastSaved.b64"))
    check("Who we are — since 1995" in xs[2] and "zero lost-time accidents" in xs[2], "the hand edit is in the file")
    check("Every lift is planned" in xs[6], "the AI-rewritten quote slide is in the file")
    check("2563EB" in z.read("ppt/theme/theme1.xml").decode(), "one tap changed the design (Ocean)")
    shot(page, "slide-7-quote", "[data-testid=slide-canvas-7]")

    # ---- 3. PDF: each slide a full 16:9 page ----
    page.locator("[data-testid=slides-save-pdf]").click(); page.wait_for_timeout(1500)
    pc = page.evaluate("window.__mock.pdfCalls.pop()")
    check(pc.get("fullPage") is True and len(pc["images"]) == 8 and pc["images"][0].startswith("data:image/jpeg"), "PDF: 8 slide pictures, full pages")
    check(page.evaluate("window.__mock.lastSaved.name") == "Adrighem & Aldibiki Cranes.pdf", "saved as a PDF")

    # ---- 4. a report from an Excel/CSV file: computed facts + a chart ----
    page.locator("[data-testid=slides-new]").click(); page.wait_for_timeout(200)
    page.locator("[data-testid=slides-tab-report]").click()
    page.fill("[data-testid=slides-prompt]", "September fleet report: revenue per crane and what to fix")
    page.locator("[data-testid=report-len-3]").click()
    csv = "Crane,Revenue EGP,Days\nLTM 1100,420000,12\nAC 500,910000,9\nGMK 3050,260000,14\nLTM 1100,180000,5\n"
    page.set_input_files("[data-testid=slides-file]", files=[{"name": "september.csv", "mimeType": "text/csv", "buffer": csv.encode()}])
    page.wait_for_selector("[data-testid=slides-file-name]", timeout=5000)
    check("4 rows" in page.locator("[data-testid=slides-file-name]").inner_text(), "the sheet is read: 4 rows")
    queue(page, [
        "TITLE: September Fleet Report\nSUBTITLE: Revenue and utilisation per crane\n1. Revenue by crane\n2. Utilisation\n3. Issues to fix",
        "The AC 500 earned the most: EGP 910,000 in 9 days.\n\nCrane | Revenue EGP\nAC 500 | 910000\nLTM 1100 | 600000\nGMK 3050 | 260000",
        "### Days on hire\n- GMK 3050: 14 days\n- LTM 1100: 17 days in two jobs",
        "- Maintenance: the GMK 3050 needs a service\n- Pricing: raise the LTM 1100 day rate by 12500",
        "The fleet is busy.\n\n1. Sales: push the AC 500 to wind-farm clients.\n2. Workshop: service the GMK 3050 this month.",
        "September revenue was EGP 1,770,000, led by the AC 500.\nFINDINGS:\n- AC 500 earned EGP 910,000\n- The LTM 1100 worked 17 days",
    ])
    page.locator("[data-testid=slides-go]").click()
    page.wait_for_selector("[data-testid=report-save-docx]", timeout=30000)
    m = sent(page)
    check(len(m) == 6, "plan + 3 sections + conclusion + summary = 6 calls (%d)" % len(m))
    check("COMPUTED FROM THE FILE" in m[1] and "total 1,770,000" in m[1], "sections get the totals computed by code from the file")
    check(page.locator("[data-testid=report-section-3]").count() == 1 and "service" in page.locator("[data-testid=report-body]").inner_text(), "the three sections are shown")
    check("1,770,000" in page.locator("[data-testid=report-summary]").inner_text(), "the executive summary is written last")
    rn = page.locator("[data-testid=report-check-note]")
    check(rn.count() == 1 and "1 figure" in rn.inner_text(), "only the invented 12500 is flagged; 600000 and 17 days are the file's own totals per crane (%s)" % rn.inner_text())
    shot(page, "04-report")
    page.locator("[data-testid=report-save-docx]").click(); page.wait_for_timeout(500)
    saved = page.evaluate("window.__mock.lastSaved")
    z = zipfile.ZipFile(io.BytesIO(base64.b64decode(saved["b64"])))
    doc = z.read("word/document.xml").decode()
    check(saved["name"] == "September Fleet Report.docx", "saved as a Word file")
    check("word/media/chart1.png" in z.namelist() and "<w:drawing>" in doc, "the Word file has the chart (a picture)")
    check('w:val="Title"' in doc and 'w:type="page"' in doc and "Executive summary" in doc and "Conclusions and recommendations" in doc, "cover title, page break, summary and conclusions")
    check(doc.index("Revenue by crane") < doc.index("Conclusions and recommendations") < doc.rindex("Data"), "sections, then conclusions, then the data")
    try:
        import docx
        d = docx.Document(io.BytesIO(base64.b64decode(saved["b64"])))
        check(len(d.inline_shapes) == 1 and len(d.tables) >= 2, "python-docx opens it: 1 chart, the tables")
    except ImportError:
        print("SKIP python-docx not installed")
    page.locator("[data-testid=report-save-pdf]").click(); page.wait_for_timeout(700)
    pb = page.evaluate("window.__mock.pdfCalls.pop().blocks")
    types = [b["type"] for b in pb]
    check(types[0] == "title" and "pagebreak" in types and "image" in types, "the PDF gets the same cover, page break and chart")

    # ---- 5. the report → a presentation, by code ----
    page.locator("[data-testid=report-to-deck]").click()
    page.wait_for_selector("[data-testid=slides-deck]", timeout=5000)
    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    z, xs = pptx_texts(page.evaluate("window.__mock.lastSaved.b64"))
    check(any("AC 500" in x and "910000" in x for x in xs) and any("Key findings" in x for x in xs), "slides from the report: key findings, the chart of the file…")
    shot(page, "05-report-deck", "[data-testid=slides-deck]")

    # ---- 6. back to the start: the recent list ----
    page.locator("[data-testid=slides-new]").click(); page.wait_for_timeout(200)
    check(page.locator("[data-testid=slides-recent]").count() == 1 and "September Fleet Report" in page.locator("[data-testid=slides-recent]").inner_text(), "recent decks and reports are listed")
    page.locator("[data-testid=slides-tab-deck]").click()
    page.fill("[data-testid=slides-prompt]", "عرض عن سوق تأجير الأوناش في مصر للمستثمرين")
    page.locator("[data-testid=slides-n-5]").click()
    page.locator("[data-testid=slides-theme-emerald]").click()
    page.locator("[data-testid=slides-web]").click()
    queue(page, [
        "TITLE: سوق تأجير الأوناش في مصر\nSUBTITLE: فرصة استثمارية\n1. [stats] السوق بالأرقام\n2. [chart] أسعار الإيجار اليومي\n3. [bullets] محركات النمو\n4. [steps] خطة الدخول\n5. [bullets] الخلاصة",
        "- 4.2 مليار جنيه | حجم السوق في 2025\n- 61% | نصيب الأوناش المتحركة\n- 1,350 | موقع نشط في العاصمة الإدارية\n- 99% | رضا العملاء\nNOTES: أرقام السوق.",
        "UNIT: جنيه في اليوم\n- 60 طن | 35,000\n- 200 طن | 90,000\n- 500 طن | 210,000\n- 800 طن | 400,000\nTAKEAWAY: الأوناش الكبيرة تحقق أعلى دخل يومي.\nNOTES: الأسعار.",
        "- العاصمة الإدارية: مشروعات بناء ضخمة\n- الطاقة: مزارع رياح جديدة\n- البنية التحتية: كباري وطرق\nNOTES: النمو.",
        "- الدراسة: تحليل الطلب\n- الأسطول: شراء أوناش مستعملة\n- المبيعات: عقود سنوية\nNOTES: الخطة.",
        "- السوق: كبير ومتنامي\n- الفرصة: الأوناش الثقيلة\n- الخطوة التالية: اجتماع\nNOTES: الخلاصة.",
    ])
    page.locator("[data-testid=slides-go]").click()
    page.wait_for_selector("[data-testid=slides-save-pptx]", timeout=30000)
    m = sent(page)
    check(page.evaluate("(window.__mock.searches || []).length") >= 2, "the web was searched (several searches in parallel)")
    check("stats (" in m[0] and "chart (" in m[0] and "Arabic" in m[0], "with sources, the plan may use numbers and charts; written in Arabic")
    check("EGP 4.2 billion" in m[1] and "copy every number exactly" in m[1], "each slide gets the passages from the pages")
    note = page.locator("[data-testid=slides-check-note]").inner_text()
    check("2 figure(s) weren't in the sources" in note, "99%% and 400,000 aren't on the pages — left out of the slides (%s)" % note)
    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    z, xs = pptx_texts(page.evaluate("window.__mock.lastSaved.b64"))
    check('rtl="1"' in xs[0] and "سوق تأجير الأوناش في مصر" in xs[0], "Arabic slides are right-to-left")
    check("شكراً لكم" in xs[-1] and "المحتويات" in xs[1], "Arabic agenda and closing slide")
    check("61%" in xs[2] and "99%" not in xs[2] and "210,000" in xs[3] and "400,000" not in xs[3], "only figures found on the pages are drawn")
    for k in (1, 2, 3, 4, 5, 6):
        shot(page, "ar-slide-%d" % k, "[data-testid=slide-canvas-%d]" % k)
    shot(page, "06-ar-deck")
    # ---- 8. from Chat: "make a presentation…" offers Slides & Reports with the prompt ----
    page.locator("nav button").first.click(); page.wait_for_timeout(300)
    queue(page, ["Here is an outline for your presentation: 1. Intro 2. Safety."] * 4)   # spare copies: any side request takes one, never the real engine
    page.fill("textarea", "Make me a presentation about tower crane safety")
    page.keyboard.press("Enter")
    try: page.wait_for_selector("[data-testid=slides-chip]", timeout=60000)
    except Exception:
        page.screenshot(path=os.environ.get("SHOTS", "/tmp") + "/chat-fail.png"); print(page.evaluate("JSON.stringify(window.__mock.bodies.map(b => [b.max_tokens, (b.messages[b.messages.length-1].content||'').slice(0,160)]))")); raise
    shot(page, "07-chat")
    page.locator("[data-testid=slides-chip]").click()
    page.wait_for_selector("[data-testid=slides-prompt]", timeout=5000)
    check(page.input_value("[data-testid=slides-prompt]") == "Make me a presentation about tower crane safety", "Chat's “Make it a PowerPoint” opens Slides & Reports with the request")
    # ---- 9. design, transitions and animations from the request itself; Play shows them ----
    page.locator("nav button").last.click(); page.wait_for_timeout(250)
    page.locator(".rounded-t-2xl button:has-text('Slides & Reports')").first.click()
    page.wait_for_selector("[data-testid=slides-prompt]", timeout=5000)
    page.locator("[data-testid=slides-tab-deck]").click()
    page.fill("[data-testid=slides-prompt]", "A pitch for our crane company. Dark blue theme, fade transitions, and the points fly in one by one")
    page.locator("[data-testid=slides-n-5]").click()
    page.locator("[data-testid=slides-theme-sunset]").click()
    queue(page, [OUTLINE] + SLIDES)
    page.locator("[data-testid=slides-go]").click()
    page.wait_for_selector("[data-testid=slides-save-pptx]", timeout=30000)
    sn = page.locator("[data-testid=slides-style-note]").inner_text()
    check("Midnight" in sn and "Fade" in sn and "Fly in" in sn, "the request's words chose the design, transition and animation (%s)" % sn)
    check("ignore any wishes about colours, design, transitions or animations" in sent(page)[0], "the AI is told to leave the design words to the app")
    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    z, xs = pptx_texts(page.evaluate("window.__mock.lastSaved.b64"))
    check(all('<p:transition spd="med"><p:fade/></p:transition>' in x for x in xs), "every slide has the fade transition")
    check(xs[2].count('presetID="2"') >= 12 and xs[2].count('clickEffect') == 3, "slide 3's three cards fly in one by one, on tap")
    check('2DD4BF' in z.read("ppt/theme/theme1.xml").decode(), "…in the Midnight (dark blue) design, not the one tapped before")
    page.select_option("[data-testid=deck-transition]", "push"); page.select_option("[data-testid=deck-trigger]", "auto")
    page.locator("[data-testid=slides-save-pptx]").click(); page.wait_for_timeout(400)
    z, xs = pptx_texts(page.evaluate("window.__mock.lastSaved.b64"))
    check('<p:push dir="u"/>' in xs[1] and "afterEffect" in xs[2] and 'evt="onBegin"' in xs[2], "changed afterwards: push transition, animations run by themselves")
    page.select_option("[data-testid=deck-trigger]", "click")
    page.locator("[data-testid=slides-play]").click()
    page.wait_for_selector("[data-testid=presenter]", timeout=5000)
    check(page.locator("[data-testid=presenter-count]").inner_text().startswith("1 / 8"), "Play starts at slide 1 of 8")
    page.wait_for_timeout(600)
    before = canvas_drawn(page, "[data-testid=presenter-canvas]")
    for _ in range(4): page.locator("[data-testid=presenter-next]").click(); page.wait_for_timeout(250)
    shot(page, "08-play")
    cnt = page.locator("[data-testid=presenter-count]").inner_text()
    check(cnt.startswith("2 / 8") and canvas_drawn(page, "[data-testid=presenter-canvas]") > 3 and before > 3, "3 taps bring in the cover's title, subtitle and date; the 4th opens slide 2 (%s)" % cnt)
    page.locator("[data-testid=presenter-close]").click(); page.wait_for_timeout(200)
    check(page.locator("[data-testid=slides-deck]").count() == 1, "closing Play goes back to the slides")
    check(not real_errors(errors), "no errors (%s)" % real_errors(errors)[:3])
    ctx.close(); br.close()
env.close()
finish()
