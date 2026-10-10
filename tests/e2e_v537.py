"""v5.37 — "as many file converters as you can": PowerPoint → PDF, PDFs merged, keep some pages,
rotate, a photo → PNG (the phone's own canvas), subtitles SRT → VTT, JSON → Excel, a web page → Word.

  python3 tests/e2e_v537.py
"""
import re, base64, io, json, os, zipfile
from playwright.sync_api import sync_playwright
from harness import Env, new_page, check, real_errors, finish, HERE

env = Env(); errors = []

# a real PowerPoint file (python-pptx), made here so this test stands on its own
from pptx import Presentation
_p = Presentation(); _s = _p.slides.add_slide(_p.slide_layouts[1]); _s.shapes.title.text = "Safety"; _s.placeholders[1].text = "Wear a helmet"
_buf = io.BytesIO(); _p.save(_buf); PPTX = _buf.getvalue()

NATIVE_PDF = """(() => { const N = window.AttuneNative, S = window.__mock;
  const R = (id, v) => setTimeout(() => window.__attuneNative.resolve(id, JSON.stringify(v)), 30);
  N.makePdf = (id, arg) => { const a = JSON.parse(arg); (S.pdfCalls = S.pdfCalls || []).push(['make', a]); R(id, { b64: btoa("%PDF-1.7 fake"), bytes: 13 }); };
  N.pdfEdit = (id, arg) => { const a = JSON.parse(arg); (S.pdfCalls = S.pdfCalls || []).push(['edit', a]);
    R(id, a.op === 'split' ? { files: [1, 2, 3].map((n) => ({ n, b64: btoa("%PDF p" + n), pages: 1 })) } : { files: [{ n: 1, b64: btoa("%PDF-1.7 edited"), pages: a.op === 'merge' ? 7 : (a.pages || []).length || 4 }] }); };
})();"""

def pick(page, name, data, mime):
    page.set_input_files("[data-testid=convert-file]", files=[{"name": name, "mimeType": mime, "buffer": data}])
def go(page):
    page.locator("[data-testid=convert-go]").click()
    page.wait_for_selector("[data-testid=convert-result]", timeout=20000)
def again(page):
    page.locator("text=Convert another file").click(); page.wait_for_timeout(150)
def last(page, kind):
    return page.evaluate("(k) => window.__mock.pdfCalls.filter(c => c[0] === k).pop()[1]", kind)
def saved(page):
    page.locator("[data-testid=convert-save]").click(); page.wait_for_timeout(250)
    return page.evaluate("window.__mock.lastSaved")

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx, page = new_page(br, env, errors)
    page.goto(env.url); page.wait_for_selector("nav", timeout=15000)
    page.evaluate(NATIVE_PDF)
    page.locator("nav button").last.click(); page.wait_for_timeout(250)
    check("80+ conversions" in page.locator(".rounded-t-2xl").inner_text(), "More says 80+ conversions")
    page.locator(".rounded-t-2xl button:has-text('File Converter')").first.click()
    page.wait_for_selector("[data-testid=convert]", timeout=5000)

    # ---- 1. PowerPoint → PDF ----
    pick(page, "safety.pptx", PPTX, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    check("PowerPoint" in page.locator("[data-testid=convert-source]").inner_text(), "a PowerPoint is recognised")
    go(page)
    mk = last(page, "make")
    check(page.locator("[data-testid=convert-out-name]").inner_text() == "safety.pdf" and [b["text"] for b in mk["blocks"]] == ["1. Safety", "Wear a helmet"], "PowerPoint → PDF: each slide becomes a heading with its text (%s)" % mk.get("blocks"))

    # ---- 2. two PDFs → merged ----
    again(page)
    page.set_input_files("[data-testid=convert-file]", files=[{"name": "a.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1 a"}, {"name": "b.pdf", "mimeType": "application/pdf", "buffer": b"%PDF-1 b"}])
    check("2 PDFs" in page.locator("[data-testid=convert-source]").inner_text() and page.locator("[data-testid=convert-to-merge]").count() == 1, "several PDFs → the merge option")
    go(page)
    e = last(page, "edit")
    check(e["op"] == "merge" and len(e["files"]) == 2 and base64.b64decode(e["files"][1]) == b"%PDF-1 b", "both PDFs go to the phone in the picked order")
    check(page.locator("[data-testid=convert-out-name]").inner_text() == "a-merged.pdf" and "7 pages" in page.locator("[data-testid=convert-result]").inner_text(), "the merged PDF: a-merged.pdf, 7 pages")

    # ---- 3. keep some pages; rotate ----
    again(page)
    pick(page, "manual.pdf", b"%PDF-1 m", "application/pdf")
    page.locator("[data-testid=convert-to-pick]").click()
    page.locator("[data-testid=convert-go]").click(); page.wait_for_timeout(200)
    check("Type the pages" in page.locator("[data-testid=convert-error]").inner_text(), "keep pages asks which pages first")
    page.fill("[data-testid=convert-pages]", "1-3, 7")
    go(page)
    check(last(page, "edit")["pages"] == [1, 2, 3, 7] and page.locator("[data-testid=convert-out-name]").inner_text() == "manual-pages.pdf", "pages 1-3 and 7 are kept → manual-pages.pdf")
    page.locator("[data-testid=convert-to-rotate]").click(); page.fill("[data-testid=convert-pages]", ""); page.locator("[data-testid=convert-deg-270]").click()
    go(page)
    e = last(page, "edit")
    check(e["op"] == "rotate" and e["degrees"] == 270 and e["pages"] == [], "rotate: all pages 90° left")
    page.locator("[data-testid=convert-to-split]").click(); go(page)
    z = zipfile.ZipFile(io.BytesIO(base64.b64decode(saved(page)["b64"])))
    check(z.namelist() == ["manual-page-01.pdf", "manual-page-02.pdf", "manual-page-03.pdf"], "split: a zip with one PDF per page")

    # ---- 4. a photo → PNG (in the page) ----
    again(page)
    jpg = page.evaluate("""() => { const c = document.createElement('canvas'); c.width = 40; c.height = 30; const g = c.getContext('2d'); g.fillStyle = '#c00'; g.fillRect(0, 0, 40, 30); return c.toDataURL('image/jpeg').split(',')[1]; }""")
    pick(page, "crane.jpg", base64.b64decode(jpg), "image/jpeg")
    page.locator("[data-testid=convert-to-png]").click(); go(page)
    s = saved(page)
    check(s["name"] == "crane.png" and base64.b64decode(s["b64"])[:4] == b"\x89PNG" and "40×30" in page.locator("[data-testid=convert-result]").inner_text(), "JPG → a real PNG file, same size (40×30)")

    # ---- 5. subtitles SRT → VTT ----
    again(page)
    pick(page, "film.srt", "1\n00:00:01,000 --> 00:00:02,500\nمرحبا\n".encode(), "application/x-subrip")
    go(page)
    s = saved(page)
    check(s["name"] == "film.vtt" and s["text"].startswith("WEBVTT") and "00:00:01.000 --> 00:00:02.500\nمرحبا" in s["text"], "SRT → VTT with Arabic")

    # ---- 6. JSON → Excel ----
    again(page)
    pick(page, "fleet.json", json.dumps([{"crane": "LTM 1100", "tons": 100}, {"crane": "ونش", "tons": 50}], ensure_ascii=False).encode(), "application/json")
    go(page)
    s = saved(page)
    xl = zipfile.ZipFile(io.BytesIO(base64.b64decode(s["b64"]))).read("xl/worksheets/sheet1.xml").decode()
    check(s["name"] == "fleet.xlsx" and "ونش" in xl and "<v>100</v>" in xl, "JSON → Excel: a column per field, numbers as numbers")

    # ---- 7. a web page → Word ----
    again(page)
    pick(page, "article.html", "<html><body><h1>Rigging</h1><p>Check every <b>sling</b>.</p><script>x()</script></body></html>".encode(), "text/html")
    page.locator("[data-testid=convert-to-docx]").click(); go(page)
    doc = zipfile.ZipFile(io.BytesIO(base64.b64decode(saved(page)["b64"]))).read("word/document.xml").decode()
    plain = re.sub(r"<[^>]+>", "", doc)
    check("Heading1" in doc and "Check every sling." in plain and "x()" not in doc, "web page → Word: the heading and text, no script")
    check(re.search(r"<w:b/>(?:<w:bCs/>)?</w:rPr><w:t[^>]*>sling</w:t>", doc) is not None, "v6.8: web page → Word: the bold word stays bold")

    # ---- 8. Arabic labels ----
    page.evaluate("localStorage.setItem('attune:ui:lang','ar')"); page.reload(); page.wait_for_selector("nav", timeout=15000)
    page.evaluate(NATIVE_PDF)
    page.locator("nav button").last.click(); page.wait_for_timeout(250)
    page.locator(".rounded-t-2xl button:has-text('محوّل الملفات')").first.click()
    page.wait_for_selector("[data-testid=convert]", timeout=5000)
    pick(page, "manual.pdf", b"%PDF-1 m", "application/pdf")
    labels = page.locator("[data-testid=convert]").inner_text()
    check("ادمجهم" not in labels and "التقسيم إلى صفحات" in labels and "الاحتفاظ بصفحات محددة" in labels and "تدوير الصفحات" in labels, "the PDF tools are in Arabic")
    check(not real_errors(errors), "no errors (%s)" % real_errors(errors)[:3])
    ctx.close(); br.close()
env.close()
finish()
