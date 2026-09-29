"""v6.8 — sample documents with every kind of formatting (Ali: "not just convert but with formatting and
spacing and everything, like from PDF to Word"). make_samples.py → tests/convert/out/*.docx (+ .pdf by LibreOffice)."""
import os, subprocess, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.path.join(os.path.dirname(__file__), "out"); os.makedirs(OUT, exist_ok=True)

def rtl(p):
    # a right-to-left paragraph: start = the right edge (no w:jc — "right" means the LEFT side in a bidi
    # paragraph), and its runs bold / sized with the complex-script tags Arabic uses (bCs, szCs)
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi"); pPr.insert(0, b)
    for jc in pPr.findall(qn("w:jc")):
        if jc.get(qn("w:val")) in ("right", "end"): pPr.remove(jc)
    for r in p.runs: cs(r)

def cs(r):
    rPr = r._r.get_or_add_rPr()
    if r.bold: rPr.append(OxmlElement("w:bCs"))
    if r.font.size:
        z = OxmlElement("w:szCs"); z.set(qn("w:val"), str(int(r.font.size.pt * 2))); rPr.append(z)

def rtl_table(tb):
    tblPr = tb._tbl.tblPr; b = OxmlElement("w:bidiVisual"); tblPr.append(b)
    for row in tb.rows:
        for c in row.cells:
            for p in c.paragraphs: rtl(p)

def picture(path, w=360, h=180):
    # a simple chart-like PNG (bars) made without extra libraries
    import zlib, struct
    rows = []
    for y in range(h):
        row = bytearray([0])
        for x in range(w):
            bar = (x // 60) % 6; top = h - (40 + bar * 22)
            c = (40, 110, 200) if (x % 60 > 8 and y > top) else (245, 245, 245)
            row += bytes(c)
        rows.append(bytes(row))
    raw = b"".join(rows)
    def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    png = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b"")
    open(path, "wb").write(png)

def report():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Liberation Serif"; st.font.size = Pt(11)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Crane Inspection Report"); r.bold = True; r.font.size = Pt(22)
    s = d.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Adrighem & Aldibiki — Site 14, New Cairo"); r.italic = True; r.font.size = Pt(12)
    d.add_heading("1. Summary", level=1)
    p = d.add_paragraph("The "); p.add_run("Liebherr LTM 1090").bold = True; p.add_run(" was inspected on "); p.add_run("12 September 2026").italic = True
    p.add_run(". All load tests passed, and the outriggers were checked at full extension. The crane is "); p.add_run("fit for service").bold = True; p.add_run(" until the next yearly check.")
    p2 = d.add_paragraph("This paragraph is justified so both edges line up. It is long enough to wrap over several lines, which shows whether the spacing between the lines and the words is kept the same as in the original document after the conversion.")
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.add_heading("2. Checks done", level=2)
    for x in ["Hydraulic hoses and fittings", "Wire rope and hook block", "Load moment indicator (LMI)"]:
        d.add_paragraph(x, style="List Bullet")
    for x in ["Extend all outriggers", "Lift the test load at 10 m", "Slew 360° slowly"]:
        d.add_paragraph(x, style="List Number")
    q = d.add_paragraph("Note: the operator must keep the logbook in the cab at all times."); q.paragraph_format.left_indent = Cm(1.5)
    f = d.add_paragraph(); fr = f.add_run("Weather on the day: wind 18 km/h, dry. The lift plan was approved by the site manager."); fr.font.name = "Liberation Sans"
    w = d.add_paragraph(); wr = w.add_run("Warning: never lift above 80% of the chart in wind over 30 km/h."); wr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); wr.bold = True
    d.add_heading("3. Results", level=2)
    tb = d.add_table(rows=4, cols=3); tb.style = "Table Grid"
    data = [["Test", "Load (t)", "Result"], ["Main boom", "25.0", "Pass"], ["Jib", "6.5", "Pass"], ["Outriggers", "—", "Pass"]]
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            c = tb.cell(i, j); c.text = v
            if i == 0: c.paragraphs[0].runs[0].bold = True
    d.add_paragraph()
    picture(os.path.join(OUT, "chart.png"))
    d.add_picture(os.path.join(OUT, "chart.png"), width=Cm(9))
    cap = d.add_paragraph("Figure 1 — Loads by boom length"); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].font.size = Pt(9)
    ar = d.add_paragraph("ملاحظة: الونش جاهز للشغل، والفحص الجاي بعد سنة."); ar.alignment = WD_ALIGN_PARAGRAPH.RIGHT; rtl(ar)
    sig = d.add_paragraph("Eng. Karim Hassan"); sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.save(os.path.join(OUT, "report.docx"))

def arabic():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "DejaVu Sans"; st.font.size = Pt(11)
    t = d.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("عرض سعر إيجار ونش"); r.bold = True; r.font.size = Pt(20); rtl(t)
    h = d.add_heading("بيانات العميل", level=2); rtl(h)
    p = d.add_paragraph()
    p.add_run("السادة / "); p.add_run("شركة أوراسكوم للإنشاءات").bold = True; p.add_run("، نتشرف بتقديم عرض السعر التالي لإيجار ونش 70 طن لمدة 5 أيام في موقع العاصمة الإدارية. ويشمل العرض التأمين على الونش طول فترة الشغل."); rtl(p)
    for x in ["السعر شامل السواق والوقود", "الدفع 50% مقدم و50% بعد انتهاء الشغل", "العرض ساري لمدة 15 يوم"]:
        li = d.add_paragraph(x, style="List Bullet"); rtl(li)
    tb = d.add_table(rows=3, cols=3); tb.style = "Table Grid"
    for i, row in enumerate([["البند", "الكمية", "السعر"], ["ونش 70 طن", "5 أيام", "75,000 جنيه"], ["نقل الونش", "1", "8,000 جنيه"]]):
        for j, v in enumerate(row):
            tb.cell(i, j).text = v
            if i == 0: tb.cell(i, j).paragraphs[0].runs[0].bold = True
    rtl_table(tb)
    d.add_paragraph()
    e = d.add_paragraph("الإجمالي: 83,000 جنيه"); e.runs[0].bold = True; rtl(e)
    sg = d.add_paragraph("م. كريم حسن"); sg.alignment = WD_ALIGN_PARAGRAPH.RIGHT; rtl(sg)   # "right" in a bidi paragraph = the left side
    d.save(os.path.join(OUT, "arabic.docx"))



def long_doc():
    """A4, Arial, a running header and page-number footer, a paragraph that goes on over a page break,
    numbered sub-headings, a two-level list, a 5-column table, a right-aligned date, a line break."""
    from docx.enum.section import WD_SECTION
    d = Document()
    sec = d.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    st = d.styles["Normal"]; st.font.name = "Liberation Sans"; st.font.size = Pt(10.5)
    hp = sec.header.paragraphs[0]; hp.text = "Adrighem & Aldibiki — Fleet Maintenance Manual"; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; hp.runs[0].font.size = Pt(8)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Page "); r.font.size = Pt(8)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = fp.add_run(); run.font.size = Pt(8)
        if kind: fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        else: it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)
    dt = d.add_paragraph("Cairo, 29 September 2026"); dt.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_heading("Fleet Maintenance Manual", level=0)
    d.add_heading("1. Daily checks", level=1)
    d.add_heading("1.1 Before starting the engine", level=2)
    body = ("Walk around the crane and look for leaks under the carrier, loose bolts on the outrigger beams and damage to the tyres. "
            "Check the engine oil, the coolant and the hydraulic tank level with the boom fully retracted, because the level reading changes when cylinders are extended. ")
    for k in range(9):
        p = d.add_paragraph(body * 2)
        if k == 3:
            p.add_run(" The operator signs the daily sheet"); br = p.add_run(); br.add_break(); p.add_run("and hands it to the site supervisor before the first lift.")
    d.add_heading("1.2 Items to check", level=2)
    for top, subs in (("Hydraulics", ["Hoses", "Cylinders and seals"]), ("Structure", ["Boom sections", "Outrigger pads"])):
        d.add_paragraph(top, style="List Bullet")
        for s in subs: d.add_paragraph(s, style="List Bullet 2")
    d.add_heading("2. Service intervals", level=1)
    tb = d.add_table(rows=5, cols=5); tb.style = "Table Grid"
    rows = [["Item", "Daily", "250 h", "500 h", "1000 h"], ["Engine oil", "Check", "Change", "Change", "Change"], ["Hydraulic filter", "—", "Check", "Change", "Change"],
            ["Slew bearing grease", "—", "Grease", "Grease", "Inspect"], ["Wire rope", "Look", "Lubricate", "Measure", "Replace if worn"]]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tb.cell(i, j).text = v
            if i == 0: tb.cell(i, j).paragraphs[0].runs[0].bold = True
    d.add_paragraph()
    picture(os.path.join(OUT, "wide.png"), 600, 160)
    pp = d.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER; pp.add_run().add_picture(os.path.join(OUT, "wide.png"), width=Cm(15))
    for k in range(3): d.add_paragraph(body)
    d.save(os.path.join(OUT, "long.docx"))

report(); arabic(); long_doc()
for f in ["report.docx", "arabic.docx", "long.docx"]:
    subprocess.run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", OUT, os.path.join(OUT, f)], check=True, capture_output=True)
print("ok", sorted(os.listdir(OUT)))
