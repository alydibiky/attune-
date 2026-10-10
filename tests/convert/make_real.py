"""v6.10 — three real-world style PDFs for the fidelity bench (fixtures/real/, with their hand-counted expect.json):
  paper_2col     a scientific paper: title block, two columns, formulas with italics, sub / superscripts and a
                 private-use symbol (OpenSymbol U+E08A, text a reader cannot trust), a small results table — printed by Chromium
  invoice_tables a table-heavy invoice: header table, 12 line items with shaded header and right-aligned money, totals — Chromium
  arabic_letter  an Arabic right-to-left letter with a few English words and numbers — Word file printed by LibreOffice
python3 make_real.py   (needs Playwright's Chromium, LibreOffice, python-docx)"""
import os, subprocess, tempfile, shutil, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "real")
HERE = os.path.dirname(os.path.abspath(__file__))

PAPER = """<!doctype html><html><head><meta charset="utf-8"><style>
@page { size: A4; margin: 2cm 1.8cm; }
body { font-family: 'Liberation Serif', 'Times New Roman', serif; font-size: 10pt; line-height: 1.25; margin: 0; }
h1 { font-size: 17pt; text-align: center; margin: 0 0 4pt; }
.auth { text-align: center; font-size: 10.5pt; margin: 0 0 2pt; } .aff { text-align: center; font-size: 9pt; font-style: italic; margin: 0 0 10pt; }
.abs { margin: 0 1.2cm 12pt; text-align: justify; font-size: 9.5pt; } .abs b { font-size: 9.5pt; }
.cols { column-count: 2; column-gap: 0.8cm; text-align: justify; }
h2 { font-size: 11pt; margin: 8pt 0 3pt; } p { margin: 0 0 5pt; }
.eq { text-align: center; font-style: italic; margin: 4pt 0 6pt; } .sym { font-family: OpenSymbol; font-style: normal; }
table { border-collapse: collapse; width: 100%; font-size: 9pt; margin: 4pt 0 6pt; } th, td { border-top: 0.6pt solid #000; border-bottom: 0.6pt solid #000; padding: 2pt 4pt; text-align: center; }
</style></head><body>
<h1>Kinetics of Substrate-Limited Growth in a Stirred Batch Reactor</h1>
<p class="auth">A. Hassan, M. Saleh and R. Farouk</p><p class="aff">Department of Biotechnology, Faculty of Science</p>
<p class="abs"><b>Abstract.</b> We measure the specific growth rate of a yeast culture on glucose and fit the Monod model to twelve batch runs. The maximum rate is 0.42 per hour and the half-saturation constant is 0.18 grams per litre. The yield of biomass on substrate stays close to 0.5 over the whole range of initial concentrations, which supports a constant-yield model for design.</p>
<div class="cols">
<h2>1. Introduction</h2>
<p>Batch cultures are the simplest way to study how cells grow on a single limiting substrate. During exponential growth the biomass concentration rises at a rate proportional to itself, and the constant of proportionality is the specific growth rate. When the substrate becomes scarce, the rate falls and the culture enters the stationary phase.</p>
<p>Monod proposed that the specific growth rate depends on the substrate concentration in the same way that an enzyme rate depends on its substrate. This simple model still describes many cultures well and is the starting point of most reactor designs.</p>
<h2>2. Model</h2>
<p>The specific growth rate follows the Monod equation:</p>
<p class="eq">μ = μ<sub>max</sub> S / (K<sub>S</sub> + S)</p>
<p>and the biomass grows exponentially while the substrate is in excess:</p>
<p class="eq">X(t) = X<sub>0</sub> e<sup>μt</sup></p>
<p>The yield coefficient links the two balances, <span class="sym">&#xE08A;</span> marking the observed value:</p>
<p class="eq">Y<sub>X/S</sub> = <span class="sym">&#xE08A;</span>X / <span class="sym">&#xE08A;</span>S</p>
<h2>3. Methods</h2>
<p>Twelve runs were made in a 2 L stirred vessel at 30 °C and pH 5.0, with initial glucose from 1 to 20 g/L. Samples were taken every hour; biomass was measured as dry weight and glucose by an enzymatic assay. Rates were found from the slope of the logarithm of biomass against time.</p>
<h2>4. Results</h2>
<p>Table 1 lists the fitted parameters. The fit is good over the whole range, and the residuals show no trend with the initial concentration.</p>
<table><tr><th>Parameter</th><th>Value</th><th>Unit</th></tr><tr><td>μ<sub>max</sub></td><td>0.42</td><td>h<sup>−1</sup></td></tr><tr><td>K<sub>S</sub></td><td>0.18</td><td>g/L</td></tr><tr><td>Y<sub>X/S</sub></td><td>0.51</td><td>g/g</td></tr></table>
<p>The doubling time at the maximum rate is ln 2 / μ<sub>max</sub> = 1.65 h, in line with earlier reports for the same strain.</p>
<h2>5. Conclusion</h2>
<p>A constant yield and Monod kinetics describe these cultures well enough for the design of a batch process.</p>
</div></body></html>"""

ROWS = [("Laptop stand, aluminium", 2, 450.00), ("USB-C hub, 7 ports", 3, 620.00), ("Wireless mouse", 5, 275.50), ("Mechanical keyboard", 2, 1850.00),
        ("27-inch monitor", 1, 7400.00), ("HDMI cable, 2 m", 6, 95.00), ("Webcam, full HD", 2, 1150.00), ("Desk lamp, LED", 3, 540.00),
        ("Office chair", 1, 4300.00), ("Noise-cancelling headset", 2, 2650.00), ("External SSD, 1 TB", 2, 2900.00), ("Surge protector", 4, 310.00)]

def invoice():
    sub = sum(q * p for _, q, p in ROWS)
    rows = "".join(f"<tr{' class=z' if i % 2 else ''}><td>{i + 1}</td><td class=l>{d}</td><td>{q}</td><td class=r>{p:,.2f}</td><td class=r>{q * p:,.2f}</td></tr>" for i, (d, q, p) in enumerate(ROWS))
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 1.8cm; }} body {{ font-family: 'Liberation Sans', Arial, sans-serif; font-size: 10pt; margin: 0; }}
h1 {{ font-size: 22pt; color: #1F4E79; margin: 0 0 6pt; }} p {{ margin: 0 0 3pt; }}
table {{ border-collapse: collapse; width: 100%; margin: 8pt 0; }} td, th {{ border: 0.6pt solid #999; padding: 3pt 5pt; text-align: center; }}
th {{ background: #1F4E79; color: #fff; }} tr.z td {{ background: #EDF2F8; }} .l {{ text-align: left; }} .r {{ text-align: right; }}
.meta td {{ border: none; text-align: left; padding: 1pt 4pt; }} .tot {{ width: 45%; margin-left: 55%; }} .tot td {{ text-align: right; }}
</style></head><body>
<h1>INVOICE</h1>
<table class="meta"><tr><td><b>Nile Office Supplies</b></td><td><b>Invoice no.</b> INV-2026-0417</td></tr><tr><td>12 Tahrir Street, Cairo</td><td><b>Date</b> 14 September 2026</td></tr><tr><td>billing@nile-office.example</td><td><b>Due</b> 14 October 2026</td></tr></table>
<p><b>Bill to:</b> Delta Engineering Consultants, 5 Corniche Road, Alexandria</p>
<table><tr><th>#</th><th>Description</th><th>Qty</th><th>Unit price (EGP)</th><th>Amount (EGP)</th></tr>{rows}</table>
<table class="tot"><tr><td>Subtotal</td><td>{sub:,.2f}</td></tr><tr><td>VAT 14%</td><td>{sub * 0.14:,.2f}</td></tr><tr><td><b>Total due</b></td><td><b>{sub * 1.14:,.2f}</b></td></tr></table>
<p>Payment by bank transfer within 30 days. Thank you for your business.</p>
</body></html>"""

def chromium(html, out):
    t = tempfile.mkdtemp(); f = os.path.join(t, "in.html"); open(f, "w").write(html)
    subprocess.run([sys.executable, os.path.join(HERE, "print_pdf.py"), f, out], check=True)

def rtl(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:bidi"); pPr.append(b)
    # (no alignment: a right-to-left paragraph starts on the right; "right" would mean its end, the left)

def arabic(out):
    d = Document(); s = d.sections[0]; s.page_width, s.page_height = Cm(21), Cm(29.7); s.left_margin = s.right_margin = Cm(2.5)
    st = d.styles["Normal"]; st.font.name = "DejaVu Sans"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:cs"), "DejaVu Sans")
    def para(t, bold=False, size=None, align=None, after=8):
        p = d.add_paragraph(); rtl(p)
        if align: p.alignment = align
        r = p.add_run(t); r.bold = bold; r.font.cs_bold = bold
        r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
        if size: r.font.size = Pt(size); r.font.cs_size = Pt(size)
        p.paragraph_format.space_after = Pt(after)
        return p
    para("شركة النيل للتوريدات", bold=True, size=16, after=2)
    para("١٢ شارع التحرير، القاهرة", after=14)
    para("التاريخ: ١٤ سبتمبر ٢٠٢٦", after=14)
    para("السيد المهندس / أحمد سليمان المحترم", bold=True)
    para("تحية طيبة وبعد،", after=10)
    para("يسعدنا أن نرسل إليكم عرض الأسعار الخاص بتجهيز مكاتب الشركة الجديدة، والذي يشمل الأجهزة والأثاث المطلوبين حسب القائمة المرفقة. وقد راعينا في هذا العرض أفضل الأسعار المتاحة مع ضمان لمدة سنتين على جميع الأجهزة.")
    para("تبلغ القيمة الإجمالية للعرض 49,500 جنيه مصري شاملة ضريبة القيمة المضافة، ويسري العرض لمدة ثلاثين يوماً من تاريخه. ويمكن الدفع على ثلاث دفعات متساوية.")
    para("نرجو التكرم بمراجعة العرض وإفادتنا بموافقتكم، ونحن على استعداد للرد على أي استفسار عبر البريد الإلكتروني sales@nile-office.example أو على الهاتف 0225551234.")
    para("وتفضلوا بقبول فائق الاحترام والتقدير،", after=24)
    para("مدير المبيعات", bold=True, after=2)
    para("محمود عبد الرحمن")
    t = tempfile.mkdtemp(); f = os.path.join(t, "arabic_letter.docx"); d.save(f)
    subprocess.run(["soffice", "-env:UserInstallation=file://" + t + "/lo", "--headless", "--convert-to", "pdf", "--outdir", t, f], check=True, capture_output=True)
    shutil.copy(os.path.join(t, "arabic_letter.pdf"), out)

chromium(PAPER, os.path.join(D, "paper_2col.pdf"))
chromium(invoice(), os.path.join(D, "invoice_tables.pdf"))
arabic(os.path.join(D, "arabic_letter.pdf"))
print("ok")
